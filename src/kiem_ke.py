"""
KIỂM KÊ VĂN BẢN
Đọc danh mục của Long (raw/Danh_muc_van_ban_phap_luat.xlsx), tìm file tương ứng trong raw/,
nhận diện định dạng gốc (Word / PDF có text / PDF scan), đếm số trang,
rồi xuất ra output/Kiem_ke_van_ban.xlsx

Cách chạy (đứng ở thư mục legal-data):
    python src\\kiem_ke.py
"""
import re
import hashlib
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import pymupdf  # đọc PDF
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent      # thư mục legal-data
RAW = ROOT / "raw"
OUT = ROOT / "output" / "Kiem_ke_van_ban.xlsx"

# Ngưỡng: trung bình mỗi trang có >= 200 ký tự chữ thì coi là PDF có text,
# ít hơn thì là PDF scan (chỉ có ảnh, chữ ký số, con dấu...)
NGUONG_KY_TU = 200


# ---------- 1. Hàm phụ ----------
def loai_van_ban(ten_file, so_hieu):
    if so_hieu.startswith("45/2019"):
        return "Bộ luật"
    if ten_file.startswith("L"):
        return "Luật"
    if ten_file.startswith("ND"):
        return "Nghị định"
    if ten_file.startswith("TT"):
        return "Thông tư"
    return "Khác"


def md5(path):
    return hashlib.md5(path.read_bytes()).hexdigest()


# Dấu hiệu trong metadata PDF (trường Creator / Producer) cho biết file được tạo bằng gì
MAY_SCAN = ["Smart Touch", "Kodak", "HP Scan", "Scanjet", "Scanner", "Universal Scan", "EZTwain", "GdPicture"]
PHAN_MEM_VAN_PHONG = ["Microsoft® Word", "Microsoft Word", "Word 20", "LibreOffice", "Writer"]


def ngay_pdf(s):
    m = re.match(r"D?:?(\d{4})(\d{2})(\d{2})", s or "")
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else None


def phan_tich_pdf(path):
    """Trả về dict: định dạng, số trang, ký tự TB/trang, metadata PDF."""
    doc = pymupdf.open(path)
    so_trang = doc.page_count
    tong_ky_tu = sum(len(trang.get_text().strip()) for trang in doc)
    tb = tong_ky_tu / so_trang if so_trang else 0
    md = doc.metadata or {}
    creator, producer = (md.get("creator") or "").strip(), (md.get("producer") or "").strip()
    nguon = creator + " " + producer
    co_chu_ky = any(w.field_type_string == "Signature" for trang in doc for w in trang.widgets())
    if tb < NGUONG_KY_TU:
        dinh_dang = "PDF scan (chụp hình)"
        phan_mem = next((k for k in MAY_SCAN if k.lower() in nguon.lower()), None)
        nguon_tao = f"Máy scan ({creator or producer})" if phan_mem else (f"Không rõ ({nguon.strip()})" if nguon.strip() else "Không có metadata")
    elif any(k.lower() in nguon.lower() for k in PHAN_MEM_VAN_PHONG):
        dinh_dang = "PDF convert từ Word"
        nguon_tao = creator or producer
    else:
        dinh_dang = "PDF có text (xuất từ phần mềm khác)"
        nguon_tao = (creator or producer) or "Không có metadata (thường là bản Công báo điện tử)"
    return dict(dinh_dang=dinh_dang, so_trang=so_trang, ky_tu=round(tb), nguon_tao=nguon_tao,
                creator=creator, producer=producer, ngay_tao=ngay_pdf(md.get("creationDate")),
                chu_ky="Có" if co_chu_ky else "Không")


def phan_tich_file(path):
    duoi = path.suffix.lower()
    if duoi == ".pdf":
        return phan_tich_pdf(path)
    if duoi in (".docx", ".doc"):
        return dict(dinh_dang=f"Word gốc ({duoi})", so_trang=None, ky_tu=None, nguon_tao="Word",
                    creator=None, producer=None, ngay_tao=None, chu_ky=None)
    return dict(dinh_dang=f"Khác ({duoi})", so_trang=None, ky_tu=None, nguon_tao=None,
                creator=None, producer=None, ngay_tao=None, chu_ky=None)


def doc_ghi_chu():
    """Đọc raw/ghichu.txt, mỗi dòng: Tên file | Ghi chú"""
    ghi_chu = {}
    f = RAW / "ghichu.txt"
    if f.exists():
        for dong in f.read_text(encoding="utf-8").splitlines()[1:]:
            if "|" in dong:
                ten, nd = dong.split("|", 1)
                ghi_chu[ten.strip()] = nd.strip()
    return ghi_chu


# ---------- 2. Tìm tất cả file trong raw/ ----------
tat_ca_file = {}
for p in RAW.rglob("*"):
    if p.is_file() and p.suffix.lower() in (".pdf", ".docx", ".doc"):
        tat_ca_file.setdefault(p.name, []).append(p)

ghi_chu = doc_ghi_chu()

# ---------- 3. Đọc danh mục của Long ----------
wb_in = openpyxl.load_workbook(RAW / "Danh_muc_van_ban_phap_luat.xlsx")
ws_in = wb_in.active
rows = []
for r in ws_in.iter_rows(min_row=2, values_only=True):
    if not r or r[0] is None:
        continue
    stt, nhom, ten, so_hieu, ten_file, link = (list(r) + [None] * 6)[:6]
    ten_file = re.sub(r"\s*\(.*\)", "", str(ten_file)).strip()   # bỏ phần "(không tải được pdf)"
    rows.append(dict(nhom=nhom, ten=ten, so_hieu=str(so_hieu), ten_file=ten_file, link=link))

# ---------- 4. Phân tích từng văn bản ----------
ket_qua = []
da_gap = {}  # md5 -> số hiệu, để phát hiện file trùng
for i, v in enumerate(rows, 1):
    stem = Path(v["ten_file"]).stem
    files = tat_ca_file.get(v["ten_file"], [])
    # văn bản bị chia nhiều phần (vd L59_2020_QH14_tiep.pdf)
    files += [p for name, ps in tat_ca_file.items() if name.startswith(stem + "_") for p in ps]

    # chỉ lấy file nằm trong đúng thư mục nhóm nếu có nhiều bản
    dung_nhom = [p for p in files if v["nhom"] and p.parent.name.split(".")[0] == str(v["nhom"]).split(".")[0]]
    files = dung_nhom or files[:1] if len(files) > 1 and not any("_tiep" in p.name for p in files) else files

    note = ghi_chu.get(v["ten_file"], "")
    if not files:
        ket_qua.append([i, v["nhom"], v["so_hieu"], loai_van_ban(v["ten_file"], v["so_hieu"]), v["ten"],
                        v["ten_file"], "", "", None, None, "THIẾU FILE", v["link"], note,
                        None, None, None, None, None])
        continue

    kq = [phan_tich_file(p) for p in files]
    dinh_dang = kq[0]["dinh_dang"]
    so_trang = sum(k["so_trang"] or 0 for k in kq)
    ky_tu = max(k["ky_tu"] or 0 for k in kq)
    meta = kq[0]

    h = md5(files[0])
    if h in da_gap and not note:
        note = f"Trùng nội dung với văn bản {da_gap[h]}"
    da_gap.setdefault(h, v["so_hieu"])

    ket_qua.append([i, v["nhom"], v["so_hieu"], loai_van_ban(v["ten_file"], v["so_hieu"]), v["ten"],
                    " + ".join(p.name for p in files),
                    str(files[0].parent.relative_to(ROOT)),
                    dinh_dang, so_trang or None, ky_tu, "Có file", v["link"], note,
                    meta["nguon_tao"], meta["creator"], meta["producer"], meta["ngay_tao"], meta["chu_ky"]])
    print(f"{i:3}. {v['ten_file']:28} {dinh_dang:22} {so_trang:5} trang")

# ---------- 5. Ghi ra Excel ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Danh muc van ban"
tieu_de = ["STT", "Nhóm", "Số hiệu", "Loại VB", "Tên / trích yếu", "Tên file", "Thư mục",
           "Định dạng gốc", "Số trang", "Ký tự TB/trang", "Tình trạng file", "Link nguồn", "Ghi chú",
           "Nguồn tạo PDF (theo metadata)", "PDF Creator", "PDF Producer", "Ngày tạo PDF", "Chữ ký số"]
ws.append(tieu_de)
for dong in ket_qua:
    ws.append(dong)

FONT = "Arial"
vien = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
mau_dd = {"PDF scan (chụp hình)": "FCE4D6", "PDF convert từ Word": "DDEBF7",
          "PDF có text (xuất từ phần mềm khác)": "E2EFDA"}
for c in ws[1]:
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=FONT, size=10)
        c.border = vien
        c.alignment = Alignment(vertical="top", wrap_text=c.column in (2, 5, 13))
    dd = row[7].value
    if dd in mau_dd:
        row[7].fill = PatternFill("solid", fgColor=mau_dd[dd])
    if row[10].value == "THIẾU FILE":
        row[10].font = Font(name=FONT, size=10, bold=True, color="C00000")
for col, w in zip("ABCDEFGHIJKLMNOPQR", [5, 20, 17, 10, 45, 30, 28, 22, 9, 11, 12, 40, 45, 34, 22, 30, 12, 10]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# Sheet tổng hợp – dùng công thức để tự cập nhật khi sửa Sheet 1
n = len(ket_qua) + 1
th = wb.create_sheet("Tong hop")
th.append(["Định dạng gốc", "Số văn bản", "Số trang"])
for dd in ["PDF convert từ Word", "PDF có text (xuất từ phần mềm khác)", "PDF scan (chụp hình)",
           "Word gốc (.docx)", "Word gốc (.doc)"]:
    r = th.max_row + 1
    th.append([dd, f"=COUNTIF('Danh muc van ban'!H2:H{n},A{r})",
               f"=SUMIF('Danh muc van ban'!H2:H{n},A{r},'Danh muc van ban'!I2:I{n})"])
r = th.max_row + 1
th.append(["Thiếu file", f"=COUNTIF('Danh muc van ban'!K2:K{n},\"THIẾU FILE\")", None])
th.append(["Tổng số dòng danh mục", f"=COUNTA('Danh muc van ban'!A2:A{n})", f"=SUM(C2:C{r - 1})"])
th.append([])
th.append([f"Cách phân loại: trung bình < {NGUONG_KY_TU} ký tự chữ/trang -> PDF scan (chụp hình), phải OCR."])
th.append(["Có text: nếu metadata Creator/Producer là Microsoft Word -> 'PDF convert từ Word', còn lại -> xuất từ phần mềm khác (vd Công báo)."])
th.append(["Metadata PDF scan ghi tên phần mềm máy scan (Kodak Smart Touch, HP Scan, IMIP Universal Scan...) -> xác nhận là bản chụp."])
th.append(["Văn bản trùng (ND141/2026, TT152/2025) xuất hiện ở 2 nhóm nên được đếm 2 lần."])
for c in th[1]:
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
for row in th.iter_rows(min_row=2):
    for c in row:
        if c.font.color is None or not c.font.bold:
            c.font = Font(name=FONT, size=10, bold=(c.row == r + 1))
th.column_dimensions["A"].width = 26
th.column_dimensions["B"].width = 12
th.column_dimensions["C"].width = 12

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"\n=> Đã ghi {len(ket_qua)} dòng vào {OUT}")
