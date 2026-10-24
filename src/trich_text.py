"""
TRÍCH TEXT TỪ VĂN BẢN (PDF có text đọc thẳng, PDF scan thì OCR)

Đầu vào : output/Kiem_ke_van_ban.xlsx (kết quả của kiem_ke.py) + các file trong raw/
Đầu ra  : output/text/<document_id>.txt   (mỗi văn bản 1 file text, có đánh dấu trang)
          output/text/_bao_cao_b7.csv     (văn bản nào đọc bằng cách nào, bao nhiêu trang)

Cách chạy (đứng ở thư mục legal-data):
    python src\\trich_text.py

LƯU Ý:
- Thư mục output/text đã có sẵn kết quả OCR, script sẽ BỎ QUA văn bản đã có file .txt.
  Chỉ khi thêm văn bản mới (vd ND337, TT133) thì nó mới OCR văn bản đó.
- Muốn OCR trên máy Windows cần cài Tesseract:
    1) Tải bộ cài: https://github.com/UB-Mannheim/tesseract/wiki
    2) Tải file vie.traineddata (tiếng Việt):
       https://github.com/tesseract-ocr/tessdata/raw/main/vie.traineddata
       chép vào C:\\Program Files\\Tesseract-OCR\\tessdata
    3) pip install pytesseract
"""
import csv
import io
import os
import time
import warnings
from multiprocessing import Pool
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl
import pymupdf

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "raw"
OUT_DIR = ROOT / "output" / "text"
KIEM_KE = ROOT / "output" / "Kiem_ke_van_ban.xlsx"

DPI = 200            # ảnh scan gốc ~200 dpi, render đúng độ phân giải này là nhanh và đủ nét
SO_LUONG_CPU = max(1, (os.cpu_count() or 2))

# Đường dẫn Tesseract (Windows). Trên Linux/Mac để None.
TESSERACT_EXE = r"C:\Program Files\Tesseract-OCR\tesseract.exe" if os.name == "nt" else None
TESSDATA_DIR = os.environ.get("TESSDATA_DIR")   # để trống thì dùng thư mục tessdata mặc định


def danh_sach_van_ban():
    """Đọc sheet "Danh muc van ban" (do kiem_ke.py tạo), lấy danh sách văn bản KHÔNG TRÙNG và có file."""
    ws = openpyxl.load_workbook(KIEM_KE)["Danh muc van ban"]
    ds, da_co = [], set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        stt, nhom, so_hieu, loai, ten, ten_file, thu_muc, dinh_dang, so_trang, _, tinh_trang, link, note = r[:13]
        if tinh_trang != "Có file" or ten_file in da_co:
            continue
        da_co.add(ten_file)
        files = [ROOT / thu_muc / f.strip() for f in ten_file.split("+")]
        doc_id = Path(files[0]).stem
        ds.append(dict(doc_id=doc_id, so_hieu=so_hieu, files=files, dinh_dang=dinh_dang, so_trang=so_trang or 0))
    return ds


def ocr_trang(trang):
    import pytesseract
    from PIL import Image
    if TESSERACT_EXE:
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_EXE
    pix = trang.get_pixmap(dpi=DPI, colorspace=pymupdf.csGRAY)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    cfg = f'--tessdata-dir "{TESSDATA_DIR}"' if TESSDATA_DIR else ""
    return pytesseract.image_to_string(img, lang="vie", config=cfg)


SOFFICE = [r"C:\Program Files\LibreOffice\program\soffice.exe", "soffice", "libreoffice"]


def doc_sang_docx(path):
    """File .doc (Word cũ) -> .docx bằng LibreOffice (cần cài LibreOffice)."""
    import shutil
    import subprocess
    import tempfile
    out_dir = Path(tempfile.mkdtemp())
    for cmd in SOFFICE:
        if Path(cmd).exists() or shutil.which(cmd):
            subprocess.run([cmd, "--headless", "--convert-to", "docx", "--outdir", str(out_dir), str(path)],
                           check=True, capture_output=True)
            return out_dir / (Path(path).stem + ".docx")
    raise RuntimeError("Cần cài LibreOffice để đọc file .doc (hoặc tự mở bằng Word rồi Save As .docx)")


def doc_word(path):
    """Đọc Word theo đúng thứ tự đoạn văn và bảng. Word không có số trang cố định -> ghi tất cả vào TRANG 1."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    path = Path(path)
    if path.suffix.lower() == ".doc":
        path = doc_sang_docx(path)
    d = docx.Document(str(path))
    dong = []
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}p"):
            s = Paragraph(el, d).text.strip()
            if s:
                dong.append(s)
        elif el.tag.endswith("}tbl"):
            for hang in Table(el, d).rows:
                o = []
                for c in hang.cells:
                    s = " ".join(c.text.split())
                    if s and (not o or o[-1] != s):          # ô gộp bị lặp -> bỏ trùng
                        o.append(s)
                if o:
                    dong.append(" | ".join(o))
    return "\n".join(dong)


def xu_ly_van_ban(vb):
    out = OUT_DIR / f"{vb['doc_id']}.txt"
    if out.exists():
        return vb["doc_id"], "đã có sẵn", vb["so_trang"], 0
    t0 = time.time()
    la_scan = "scan" in (vb["dinh_dang"] or "")
    phan, so_trang = [], 0
    la_word = str(vb["files"][0]).lower().endswith((".doc", ".docx"))
    if la_word:
        so_trang = 1
        phan.append(f"<<<TRANG 1>>>\n{doc_word(vb['files'][0]).strip()}\n")
    for f in ([] if la_word else vb["files"]):
        doc = pymupdf.open(f)
        for trang in doc:
            so_trang += 1
            text = ocr_trang(trang) if la_scan else trang.get_text()
            phan.append(f"<<<TRANG {so_trang}>>>\n{text.strip()}\n")
    tmp = out.with_suffix(".tmp")
    tmp.write_text("\n".join(phan), encoding="utf-8")
    tmp.replace(out)  # ghi xong mới đổi tên, tránh file dở dang nếu bị ngắt giữa chừng
    cach = "Đọc file Word (python-docx)" if la_word else ("OCR (Tesseract vie)" if la_scan else "Đọc text trực tiếp")
    giay = round(time.time() - t0)
    print(f"  xong {vb['doc_id']:22} {cach:22} {so_trang:4} trang  {giay:5}s", flush=True)
    return vb["doc_id"], cach, so_trang, giay


if __name__ == "__main__":
    os.environ.setdefault("OMP_THREAD_LIMIT", "1")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ds = sorted(danh_sach_van_ban(), key=lambda v: -v["so_trang"])  # văn bản dài chạy trước
    print(f"Có {len(ds)} văn bản. Chạy song song {SO_LUONG_CPU} tiến trình...")
    with Pool(SO_LUONG_CPU) as pool:
        ket_qua = pool.map(xu_ly_van_ban, ds, chunksize=1)

    bao_cao = OUT_DIR / "_bao_cao_b7.csv"
    cu = {}
    if bao_cao.exists():
        with open(bao_cao, encoding="utf-8-sig") as f:
            cu = {r["document_id"]: r for r in csv.DictReader(f)}
    with open(bao_cao, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["document_id", "cach_doc", "so_trang", "thoi_gian_giay"])
        for doc_id, cach, st, giay in sorted(ket_qua):
            if cach == "đã có sẵn" and doc_id in cu:
                r = cu[doc_id]
                w.writerow([doc_id, r["cach_doc"], r["so_trang"], r["thoi_gian_giay"]])
            else:
                w.writerow([doc_id, cach, st, giay])
    print(f"\n=> Xong. Text nằm trong {OUT_DIR}")
