"""
CHUNKING: 2 PHƯƠNG ÁN VÀ SO SÁNH

Phương án A – Legal-structure chunking (đề xuất):
    1. Điều (kể cả header ngữ cảnh) <= 512 token  -> 1 Điều = 1 chunk
    2. Điều dài hơn                                -> cắt theo Khoản
    3. Khoản vẫn dài hơn                           -> cắt theo Điểm (kèm câu dẫn của Khoản)
    4. Điểm vẫn dài hơn                            -> cắt theo câu/đoạn, đánh dấu "split"
    5. Phần mở đầu, Phụ lục, bảng biểu             -> chunk riêng, gắn chunk_type tương ứng
    Mỗi chunk có context_header: [Loại VB số hiệu > Chương > Điều > Khoản > Điểm]

Phương án B – Fixed-token chunking (để so sánh):
    Nối toàn văn, cắt cứng mỗi 512 token, chồng lấn (overlap) 50 token.

Đầu vào : output/structure/*.json (kết quả của tach_van_ban.py)
Đầu ra  : output/chunks/legal_structure_chunks.jsonl
          output/chunks/fixed_token_chunks.jsonl
          Sheet "So sanh chunking" và "De xuat chunking" trong output/Kiem_ke_van_ban.xlsx

Cách chạy (đứng ở thư mục legal-data):
    python src\\chunking.py
"""
import json
import re
import unicodedata
import statistics
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from regex_luat import dem_token, doc_tokenizer, giai_ma, ma_hoa
from tach_van_ban import dinh_dang_sheet, phan_vi

ROOT = Path(__file__).resolve().parent.parent
STRUCT_DIR = ROOT / "output" / "structure"
CHUNK_DIR = ROOT / "output" / "chunks"
TEXT_DIR = ROOT / "output" / "text"
KIEM_KE = ROOT / "output" / "Kiem_ke_van_ban.xlsx"

MAX_TOKEN = 512      # giới hạn 1 chunk (bge-m3 nhận tới 8192, nhưng chunk ngắn truy xuất chính xác hơn)
OVERLAP = 50         # chồng lấn cho phương án fixed-token
CAU_DAN_TOI_DA = 120  # câu dẫn của Khoản ngắn hơn ngưỡng này thì chép kèm vào từng Điểm


# ============================ PHƯƠNG ÁN A ============================
def header(doc, u=None, dieu=None, khoan=None, diem=None, phu=None):
    phan = [f"{doc['document_type']} {doc['document_no']}"]
    if u is not None and u.get("chapter"):
        ten = f" {u['chapter_title']}" if u.get("chapter_title") else ""
        phan.append(f"Chương {u['chapter']}{ten}")
    if u is not None and u.get("section"):
        phan.append(f"Mục {u['section']}")
    if dieu:
        phan.append(dieu)
    if khoan:
        phan.append(f"Khoản {khoan}")
    if diem:
        phan.append(f"Điểm {diem}")
    if phu:
        phan.append(phu)
    return "[" + " > ".join(phan) + "]"


def cat_nho(text, ngan_sach):
    """Cắt đoạn text quá dài thành các phần <= ngan_sach token, ưu tiên cắt ở ranh giới dòng/câu."""
    manh = [m for m in re.split(r"(?<=[\.;:])\s+|\n", text) if m.strip()]
    phan, cur = [], ""
    for m in manh:
        thu = (cur + "\n" + m).strip()          # giữ mỗi câu một dòng để đối chiếu nguyên văn
        if dem_token(thu) <= ngan_sach:
            cur = thu
            continue
        if cur:
            phan.append(cur)
        if dem_token(m) <= ngan_sach:
            cur = m
        else:                                           # 1 câu quá dài: cắt cứng theo token
            ids = ma_hoa(m)
            for s in range(0, len(ids), ngan_sach):
                phan.append(giai_ma(ids[s:s + ngan_sach]))
            cur = ""
    if cur:
        phan.append(cur)
    return phan


def legal_chunks(doc):
    units = doc["legal_units"]
    by_id = {u["unit_id"]: u for u in units}
    con = {}
    for u in units:
        if u.get("parent_id"):
            con.setdefault(u["parent_id"], []).append(u)

    def nhan(u):
        return {"clause": f"{u['clause']}. ", "point": f"{u['point']}) "}.get(u["unit_type"], "")

    def toan_van(u, co_nhan=False):
        dau = (nhan(u) if co_nhan else "") + u["content"]
        phan = [dau] + [toan_van(c, True) for c in con.get(u["unit_id"], [])]
        return "\n".join(p for p in phan if p.strip())

    chunks = []

    def them(loai, unit, hdr, noi_dung, unit_ids, **kw):
        noi_dung = noi_dung.strip()
        if not noi_dung:
            return
        ngan_sach = MAX_TOKEN - dem_token(hdr) - 1
        phan = [noi_dung] if dem_token(noi_dung) <= ngan_sach else cat_nho(noi_dung, ngan_sach)
        for i, p in enumerate(phan, 1):
            cid = unit["unit_id"] + (f"_part{i}" if len(phan) > 1 else "")
            chunks.append({
                "chunk_id": cid,
                "strategy": "legal_structure",
                "chunk_type": loai if len(phan) == 1 else f"{loai}_split",
                "document_id": doc["document_id"],
                "document_no": doc["document_no"],
                "document_type": doc["document_type"],
                "chapter": unit.get("chapter"),
                "section": unit.get("section"),
                "article": unit.get("article"),
                "clause": kw.get("clause", unit.get("clause")),
                "point": kw.get("point", unit.get("point")),
                "unit_ids": unit_ids,
                "parent_id": kw.get("parent_id"),
                "context_header": hdr,
                "content": p,
                "text_for_embedding": hdr + "\n" + p,
                "token_count": dem_token(hdr + "\n" + p),
                "page": unit.get("page"),
                "page_end": unit.get("page"),
                "has_table": unit.get("has_table", False),
                # hiệu lực lấy theo unit (có Điều hiệu lực riêng khác toàn văn bản)
                "effective_from": unit.get("effective_from", doc["effective_from"]),
                "effective_to": unit.get("effective_to", doc["effective_to"]),
                "status": unit.get("status", doc["status"]),
                "is_confident": all(by_id[x].get("is_confident", True) for x in unit_ids if x in by_id),
                # metadata văn bản – để WP2/WP4/WP5 lọc và trích dẫn mà không cần tra lại file văn bản
                "document_title": doc.get("title"),
                "issuer": doc.get("issuer"),
                "issued_date": doc.get("issued_date"),
                "amended_by": doc.get("amended_by", []),
                "replaced_by": doc.get("replaced_by", []),
                "source_file": (doc.get("source_file") or [None])[0],
                "source_url": doc.get("source_url"),
                "extraction_method": (doc.get("extraction") or {}).get("method"),
            })

    for u in units:
        t = u["unit_type"]
        if t == "preamble":
            them("preamble", u, header(doc, phu="Phần mở đầu"), u["content"], [u["unit_id"]])
        elif t == "appendix":
            loai = "table" if u.get("has_table") else "appendix"
            them(loai, u, header(doc, phu=u["title"][:80]), u["content"], [u["unit_id"]])
        elif t == "article":
            ten_dieu = u["title"]
            hdr = header(doc, u, ten_dieu)
            day_du = toan_van(u)
            if dem_token(hdr + "\n" + day_du) <= MAX_TOKEN:                       # Quy tắc 1
                ids = [u["unit_id"]] + [x["unit_id"] for x in units if x["unit_id"].startswith(u["unit_id"] + "_")]
                them("article", u, hdr, day_du, ids)
                continue
            # Quy tắc 2: cắt theo Khoản (Điểm nằm thẳng dưới Điều cũng xử lý như Khoản)
            if u["content"]:
                them("article_intro", u, hdr, u["content"], [u["unit_id"]])
            for k in con.get(u["unit_id"], []):
                nhan_k = k["clause"]
                hdr_k = header(doc, u, ten_dieu, khoan=nhan_k) if k["unit_type"] == "clause" \
                    else header(doc, u, ten_dieu, diem=k["point"])
                full_k = toan_van(k, True)
                if dem_token(hdr_k + "\n" + full_k) <= MAX_TOKEN:
                    ids = [k["unit_id"]] + [x["unit_id"] for x in con.get(k["unit_id"], [])]
                    them(k["unit_type"], k, hdr_k, full_k, ids, parent_id=u["unit_id"])
                    continue
                # Quy tắc 3: cắt theo Điểm
                if not con.get(k["unit_id"]):                 # Khoản không có Điểm -> quy tắc 4 (cắt câu)
                    them(k["unit_type"], k, hdr_k, nhan(k) + k["content"], [k["unit_id"]], parent_id=u["unit_id"])
                    continue
                cau_dan = nhan(k) + k["content"] if k["content"] else ""
                kem = cau_dan if dem_token(cau_dan) <= CAU_DAN_TOI_DA else ""
                if cau_dan and not kem:
                    them("clause_intro", k, hdr_k, cau_dan, [k["unit_id"]], parent_id=u["unit_id"])
                for d in con.get(k["unit_id"], []):
                    hdr_d = header(doc, u, ten_dieu, khoan=nhan_k, diem=d["point"])
                    noi = (kem + "\n" + f"{d['point']}) " + d["content"]) if kem else f"{d['point']}) " + d["content"]
                    them("point", d, hdr_d, noi, [d["unit_id"]], parent_id=k["unit_id"])
    gan_trang(chunks, doc["document_id"], {u["unit_id"]: u.get("page") for u in units})
    return chunks


# ============================ XÁC ĐỊNH TRANG ============================
def chuan(s):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", s or "")).strip()


def trang_van_ban(doc_id):
    raw = (TEXT_DIR / f"{doc_id}.txt").read_text(encoding="utf-8")
    return {int(so): chuan(nd) for so, nd in re.findall(r"<<<TRANG (\d+)>>>\n(.*?)(?=<<<TRANG \d+>>>|\Z)", raw, re.S)}


def _bo_nhan(s):
    return re.sub(r"^(\d{1,2}\s*[\.,]|[a-zđ]\s*\))\s*", "", chuan(s))


def tim_trang(doan, trang, tu):
    """Trang đầu tiên (tính từ trang `tu`-1) có chứa đoạn văn `doan`."""
    if not doan:
        return None
    # dùng đoạn dài (70 ký tự) để tránh nhầm với câu khác có cùng mở đầu ("Trường hợp hộ kinh doanh, ...")
    mau = [doan[:70], doan[10:60], doan[-50:]]
    for p in sorted(q for q in trang if q >= (tu or 1)):
        if any(len(m) >= 25 and m in trang[p] for m in mau):
            return p
    for p in sorted(q for q in trang if q >= (tu or 1) - 1):      # dự phòng: đoạn ngắn hơn
        if any(len(m) >= 12 and m in trang[p] for m in (doan[:35], doan[5:30])):
            return p
    return None


def gan_trang(chunks, doc_id, trang_unit):
    """Ghi page (trang bắt đầu) và page_end (trang kết thúc) THẬT của từng chunk bằng cách tìm nội dung
    chunk trong text từng trang – quan trọng với chunk bị cắt nhỏ (_partN) nằm ở trang sau."""
    trang = trang_van_ban(doc_id)
    for c in chunks:
        dong = [_bo_nhan(x) for x in c["content"].split("\n")]
        dong = [x for x in dong if len(x) >= 12]
        if not dong:
            continue
        tu = trang_unit.get(c["unit_ids"][0]) or 1
        # Điểm có chép kèm câu dẫn của Khoản: trang bắt đầu tính theo dòng của chính Điểm
        dau = dong[1] if c["chunk_type"].startswith("point") and len(dong) > 1 and c.get("point") and \
            c["content"].split("\n")[1].startswith(f"{c['point']})") else dong[0]
        p = tim_trang(dau, trang, tu)
        c["page"] = p or c["page"]
        q = tim_trang(dong[-1], trang, c["page"])
        c["page_end"] = max(q, c["page"]) if q and c["page"] else c["page"]


# ============================ PHƯƠNG ÁN B ============================
def fixed_chunks(doc):
    """Nối toàn văn theo thứ tự, cắt cứng 512 token, overlap 50.
    Ghi lại mỗi chunk phủ những Điều nào để đo việc cắt ngang Điều."""
    ids, nhan = [], []          # nhan[i] = số Điều của token i (None = mở đầu / phụ lục)
    ranh_gioi = {0}             # vị trí token bắt đầu mỗi Điều / Khoản / Điểm
    for u in doc["legal_units"]:
        tieu_de = u["title"] if u["unit_type"] in ("article", "appendix") else None
        nhan_dau = {"point": f"{u['point']}) ", "clause": f"{u['clause']}. "}.get(u["unit_type"], "")
        text = "\n".join(x for x in [tieu_de, nhan_dau + u["content"]] if x and x.strip())
        t = ma_hoa(text + "\n")
        ranh_gioi.add(len(ids))
        ids += t
        nhan += [u["article"]] * len(t)
    chunks, buoc, i, stt = [], MAX_TOKEN - OVERLAP, 0, 0
    while i < len(ids):
        j = min(i + MAX_TOKEN, len(ids))
        dieu = sorted({a for a in nhan[i:j] if a is not None})
        dau_dieu = i in ranh_gioi          # chunk có bắt đầu đúng chỗ bắt đầu 1 Điều/Khoản/Điểm không?
        noi_dung = giai_ma(ids[i:j])
        stt += 1
        chunks.append({
            "chunk_id": f"{doc['document_id']}_F{stt:04d}",
            "strategy": "fixed_token",
            "document_id": doc["document_id"],
            "document_no": doc["document_no"],
            "document_type": doc["document_type"],
            "chunk_index": stt,
            "token_start": i,
            "token_end": j,
            "articles_covered": dieu,
            "starts_at_unit_boundary": dau_dieu,
            "content": noi_dung,
            "text_for_embedding": noi_dung,
            "token_count": j - i,
            "effective_from": doc["effective_from"],
            "effective_to": doc["effective_to"],
            "status": doc["status"],
            "document_title": doc.get("title"),
            "source_file": (doc.get("source_file") or [None])[0],
        })
        if j == len(ids):
            break
        i += buoc
    return chunks


# ============================ SO SÁNH ============================
NGOAI_THAN = ("preamble", "appendix", "table")


def chi_so(chunks, loai):
    """Các chỉ số so sánh. Phần % tính trên PHẦN THÂN văn bản (Điều/Khoản/Điểm),
    không tính phần mở đầu và phụ lục/bảng biểu (vì phụ lục không có cấu trúc Điều để giữ)."""
    tok = [c["token_count"] for c in chunks]
    if loai == "legal":
        than = [c for c in chunks if c["chunk_type"].replace("_split", "") not in NGOAI_THAN]
        cat_ngang = sum(1 for c in than if c["chunk_type"].endswith("_split"))
        nhieu_dieu = 0
        dung_1 = sum(1 for c in than if c["article"] is not None)
    else:
        than = [c for c in chunks if c["articles_covered"]]
        cat_ngang = sum(1 for c in than if not c["starts_at_unit_boundary"])
        nhieu_dieu = sum(1 for c in than if len(c["articles_covered"]) > 1)
        dung_1 = sum(1 for c in than if len(c["articles_covered"]) == 1)
    n = len(than) or 1
    return {
        "Số chunk (toàn bộ)": len(chunks),
        "  trong đó: chunk phần thân (Điều/Khoản/Điểm)": len(than),
        "  trong đó: chunk mở đầu + phụ lục/bảng biểu": len(chunks) - len(than),
        "Tổng token (kể cả header / overlap)": sum(tok),
        "Token TB / chunk": round(statistics.mean(tok)) if tok else 0,
        "Trung vị (P50)": phan_vi(tok, 50),
        "P95": phan_vi(tok, 95),
        "Lớn nhất": max(tok, default=0),
        "Nhỏ nhất": min(tok, default=0),
        f"% chunk > {MAX_TOKEN} token": round(sum(t > MAX_TOKEN for t in tok) / max(len(tok), 1), 4),
        "% chunk phần thân bị cắt ngang giữa chừng (không bắt đầu ở ranh giới Điều/Khoản/Điểm)": round(cat_ngang / n, 4),
        "% chunk phần thân chứa nội dung của nhiều hơn 1 Điều": round(nhieu_dieu / n, 4),
        "% chunk phần thân gắn được với đúng 1 Điều (trích dẫn chính xác)": round(dung_1 / n, 4),
    }


def ghi_excel(ma, mb, theo_loai, mau):
    wb = openpyxl.load_workbook(KIEM_KE)
    for ten in ("So sanh chunking", "De xuat chunking"):
        if ten in wb.sheetnames:
            del wb[ten]

    ws = wb.create_sheet("So sanh chunking")
    ws.append(["Chỉ số", "A. Legal-structure (đề xuất)", f"B. Fixed-token ({MAX_TOKEN}/{OVERLAP})"])
    for k in ma:
        ws.append([k, ma[k], mb[k]])
        if k.startswith("%"):
            ws.cell(ws.max_row, 2).number_format = ws.cell(ws.max_row, 3).number_format = "0.0%"
    ws.append([])
    ws.append(["Số chunk phương án A theo loại", "Số chunk", "Giải thích"])
    giai_thich = {
        "article": "Cả Điều vừa 1 chunk (quy tắc 1)", "clause": "Điều dài, cắt theo Khoản (quy tắc 2)",
        "point": "Khoản dài, cắt theo Điểm (quy tắc 3)", "article_intro": "Câu dẫn của Điều dài",
        "clause_intro": "Câu dẫn dài của Khoản", "preamble": "Phần mở đầu / căn cứ ban hành",
        "appendix": "Phụ lục, biểu mẫu", "table": "Phụ lục có bảng biểu",
    }
    hang_dau = ws.max_row
    for k, v in sorted(theo_loai.items(), key=lambda x: -x[1]):
        goc = k.replace("_split", "")
        ws.append([k, v, giai_thich.get(goc, "") + (" – quá dài, cắt nhỏ thêm (quy tắc 4)" if k.endswith("_split") else "")])
    for c in ws[hang_dau]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1F4E78")
    ws.append([])
    ws.append([f"Token đếm bằng: {doc_tokenizer()[0]}. Token của phương án A tính cả context_header."])
    dinh_dang_sheet(ws, [80, 26, 26])

    # ----- Sheet đề xuất -----
    dx = wb.create_sheet("De xuat chunking")
    dx.append(["Hạng mục", "Nội dung đề xuất"])
    dong = [
        ("Ngưỡng token", f"{MAX_TOKEN} token / chunk (tính cả header ngữ cảnh), đếm bằng tokenizer của bge-m3"),
        ("Quy tắc 1", "Điều ≤ ngưỡng → 1 Điều = 1 chunk"),
        ("Quy tắc 2", "Điều > ngưỡng → cắt theo Khoản (câu dẫn của Điều thành chunk article_intro)"),
        ("Quy tắc 3", f"Khoản > ngưỡng → cắt theo Điểm; câu dẫn của Khoản ≤ {CAU_DAN_TOI_DA} token được chép kèm vào từng Điểm"),
        ("Quy tắc 4", "Điểm vẫn > ngưỡng → cắt theo câu/dòng, chunk_type thêm hậu tố _split, chunk_id thêm _partN"),
        ("Quy tắc 5", "Phần mở đầu, Phụ lục, biểu mẫu, bảng → chunk riêng: preamble / appendix / table"),
        ("Header ngữ cảnh", "Mỗi chunk có context_header, ví dụ: [Nghị định 254/2026/NĐ-CP > Chương II > Điều 9. Thời điểm lập hóa đơn > Khoản 2]. "
                            "Văn bản đem embedding = context_header + nội dung"),
        ("Mã định danh", "chunk_id = unit_id theo WP1: <document_id>_D<điều>_K<khoản>_P<điểm>, ví dụ ND254_2026_D9_K2_Pa"),
        ("Metadata thời gian", "Mỗi chunk mang effective_from / effective_to / status để WP4 lọc theo ngày hiệu lực"),
        ("Độ tin cậy", "is_confident = false khi văn bản OCR có nhiều ký tự lạ hoặc số Điều bị nhảy → cần người kiểm duyệt"),
        ("Phương án so sánh", f"Fixed-token: cắt cứng {MAX_TOKEN} token, overlap {OVERLAP} token (xem sheet So sanh chunking)"),
        ("File kết quả", "output/chunks/legal_structure_chunks.jsonl và output/chunks/fixed_token_chunks.jsonl (mỗi dòng 1 chunk JSON)"),
    ]
    for d in dong:
        dx.append(list(d))
    dx.append([])
    dx.append(["Ví dụ 1 chunk (phương án A)", json.dumps({k: v for k, v in mau.items() if k != "text_for_embedding"},
                                                        ensure_ascii=False, indent=2)])
    dinh_dang_sheet(dx, [26, 110])
    for row in dx.iter_rows(min_row=2):
        row[0].font = Font(name="Arial", bold=True, size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(KIEM_KE)


def main():
    CHUNK_DIR.mkdir(parents=True, exist_ok=True)
    print("Tokenizer:", doc_tokenizer()[0])
    tat_ca_a, tat_ca_b = [], []
    for f in sorted(STRUCT_DIR.glob("*.json")):
        doc = json.loads(f.read_text(encoding="utf-8"))
        a, b = legal_chunks(doc), fixed_chunks(doc)
        tat_ca_a += a
        tat_ca_b += b
        print(f"  {doc['document_id']:22} A: {len(a):5} chunk   B: {len(b):5} chunk")

    for ten, ds in [("legal_structure_chunks.jsonl", tat_ca_a), ("fixed_token_chunks.jsonl", tat_ca_b)]:
        with open(CHUNK_DIR / ten, "w", encoding="utf-8") as out:
            for c in ds:
                out.write(json.dumps(c, ensure_ascii=False) + "\n")

    ma, mb = chi_so(tat_ca_a, "legal"), chi_so(tat_ca_b, "fixed")
    theo_loai = {}
    for c in tat_ca_a:
        theo_loai[c["chunk_type"]] = theo_loai.get(c["chunk_type"], 0) + 1
    mau = next((c for c in tat_ca_a if c["chunk_type"] == "clause" and c["document_id"].startswith("ND254")),
               next((c for c in tat_ca_a if c["chunk_type"] == "clause"), tat_ca_a[0]))
    ghi_excel(ma, mb, theo_loai, mau)

    print(f"\n{'Chỉ số':88} {'A. Legal':>9} {'B. Fixed':>9}")
    for k in ma:
        print(f"{k:88} {ma[k]:>9} {mb[k]:>9}")
    print(f"\n=> Chunk lưu trong {CHUNK_DIR}")


if __name__ == "__main__":
    main()
